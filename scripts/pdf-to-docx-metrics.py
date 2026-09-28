"""Metrics for the PDF -> Word converter check (driven by scripts/pdf-to-docx-check.mjs).

  python scripts/pdf-to-docx-metrics.py <job.json>              scores one document, writes <outDir>/metrics.json
  python scripts/pdf-to-docx-metrics.py --selftest <gt.docx> <gt.pdf>
                                  checks the .docx style resolver against Word's own rendering of the same file
  python scripts/pdf-to-docx-metrics.py --dump <file.docx>      paragraphs, styles and tables as Word would see them

Job: {"id", "orig": pdf, "gt": docx|null, "candidates": {"ours": {"docx", "pdf"|null}, "word": {...}},
      "outDir", "sheets": {"dpi", "maxPages", "overlay"}|null, "visualDpi"}

Every candidate is scored against the original PDF (text similarity, reading order, rendered style agreement,
visual difference, structure counts) and, when a ground-truth .docx exists, against it (style agreement per word,
paragraph breaks and properties, table cells, images). Rendered = the candidate .docx exported to PDF by Word.
"""

import difflib
import json
import os
import re
import sys
import time
import unicodedata
from collections import Counter
from bisect import bisect_left

import fitz
import numpy as np
from docx import Document
from PIL import Image, ImageDraw, ImageFont

NS = {
    "w": "http://schemas.openxmlformats.org/wordprocessingml/2006/main",
    "mc": "http://schemas.openxmlformats.org/markup-compatibility/2006",
    "wp": "http://schemas.openxmlformats.org/drawingml/2006/wordprocessingDrawing",
    "a": "http://schemas.openxmlformats.org/drawingml/2006/main",
    "pic": "http://schemas.openxmlformats.org/drawingml/2006/picture",
    "wps": "http://schemas.microsoft.com/office/word/2010/wordprocessingShape",
    "v": "urn:schemas-microsoft-com:vml",
    "r": "http://schemas.openxmlformats.org/officeDocument/2006/relationships",
}


def q(name):
    prefix, local = name.split(":")
    return "{%s}%s" % (NS[prefix], local)


W_P, W_R, W_T, W_TBL, W_TR, W_TC = q("w:p"), q("w:r"), q("w:t"), q("w:tbl"), q("w:tr"), q("w:tc")
W_TAB, W_BR, W_CR, W_SYM = q("w:tab"), q("w:br"), q("w:cr"), q("w:sym")
W_PPR, W_RPR, W_SDT, W_SDTCONTENT = q("w:pPr"), q("w:rPr"), q("w:sdt"), q("w:sdtContent")
W_TXBX = q("w:txbxContent")
MC_ALT, MC_CHOICE, MC_FALLBACK = q("mc:AlternateContent"), q("mc:Choice"), q("mc:Fallback")
W_VAL = q("w:val")
SKIP_TEXT = {q("w:instrText"), q("w:delText"), q("w:delInstrText"), q("w:pPr"), q("w:rPr"), q("w:tblPr"), q("w:tcPr"), q("w:trPr")}

# ── Text normalisation ─────────────────────────────────────────────────────────────────────────────────────────────

BULLETS = "\u2022\u25aa\u25cf\u25e6\u2023\u2043\u2219\u25a0\u25a1\u25c6\u25c7\u2756\u27a2\u2713\u2714\uf0b7\uf0a7\uf0d8\uf076\uf0fc\uf06e\uf0a8\uf0de\uf0e0"
BULLET_RE = re.compile("[" + BULLETS + "]")
ZERO_WIDTH = re.compile("[\u00ad\u200b\u200c\u200d\u2060\ufeff]")


def norm_text(s):
    s = unicodedata.normalize("NFKC", s or "")
    s = ZERO_WIDTH.sub("", s)
    s = BULLET_RE.sub("\u2022", s)
    s = s.replace("\u2010", "-").replace("\u2011", "-")
    s = re.sub(r"(\w)-[ \t]*\n[ \t]*(?=[a-záéíóúñü])", r"\1", s)  # words split by a line-end hyphen
    return re.sub(r"\s+", " ", s).strip()


def tokens_of(text):
    """Word tokens for alignment: lower-case letters and digits."""
    return re.findall(r"\w+", norm_text(text).lower())


# ── Sequence algorithms ────────────────────────────────────────────────────────────────────────────────────────────


def levenshtein(a, b):
    """Global edit distance (insert/delete/substitute), bit-parallel (Myers 1999 / Hyyrö)."""
    if a == b:
        return 0
    if not a:
        return len(b)
    if not b:
        return len(a)
    if len(a) > len(b):
        a, b = b, a
    m = len(a)
    peq = {}
    for i, ch in enumerate(a):
        peq[ch] = peq.get(ch, 0) | (1 << i)
    mask = (1 << m) - 1
    last = 1 << (m - 1)
    pv, mv, score = mask, 0, m
    for ch in b:
        eq = peq.get(ch, 0)
        xv = eq | mv
        xh = (((eq & pv) + pv) ^ pv) | eq
        ph = mv | (~(xh | pv) & mask)
        mh = pv & xh
        if ph & last:
            score += 1
        elif mh & last:
            score -= 1
        ph = ((ph << 1) | 1) & mask
        mh = (mh << 1) & mask
        pv = mh | (~(xv | ph) & mask)
        mv = ph & xv
    return score


def text_similarity(a, b):
    """1 - normalised edit distance of whitespace-normalised texts (None when both are empty)."""
    a, b = norm_text(a), norm_text(b)
    if not a and not b:
        return None
    if max(len(a), len(b)) > 400_000:  # too long for the exact distance: word-level approximation
        ta, tb = a.split(" "), b.split(" ")
        matched = sum(len(ta[i]) + 1 for i, _ in align_tokens(ta, tb))
        return round(2 * matched / (len(a) + len(b) + 2), 4)
    return round(1 - levenshtein(a, b) / max(len(a), len(b)), 4)


def bag_f1(ref_tokens, cand_tokens):
    """Order-insensitive word completeness: F1 of the two word multisets."""
    if not ref_tokens and not cand_tokens:
        return None
    common = sum((Counter(ref_tokens) & Counter(cand_tokens)).values())
    return round(2 * common / (len(ref_tokens) + len(cand_tokens)), 4)


def lis_indices(seq):
    """Indices of a longest strictly increasing subsequence of seq."""
    tails, tails_idx, prev = [], [], [-1] * len(seq)
    for i, x in enumerate(seq):
        k = bisect_left(tails, x)
        if k == len(tails):
            tails.append(x)
            tails_idx.append(i)
        else:
            tails[k] = x
            tails_idx[k] = i
        prev[i] = tails_idx[k - 1] if k else -1
    out, i = [], tails_idx[-1] if tails_idx else -1
    while i >= 0:
        out.append(i)
        i = prev[i]
    return out[::-1]


def align_tokens(a, b):
    """Matched (i, j) pairs, increasing in both, between token lists (patience diff, difflib inside small gaps)."""
    pairs = []
    stack = [(0, len(a), 0, len(b))]
    while stack:
        a0, a1, b0, b1 = stack.pop()
        while a0 < a1 and b0 < b1 and a[a0] == b[b0]:
            pairs.append((a0, b0))
            a0 += 1
            b0 += 1
        while a0 < a1 and b0 < b1 and a[a1 - 1] == b[b1 - 1]:
            a1 -= 1
            b1 -= 1
            pairs.append((a1, b1))
        if a0 >= a1 or b0 >= b1:
            continue
        ca, cb = Counter(a[a0:a1]), Counter(b[b0:b1])
        pos_b = {b[j]: j for j in range(b0, b1) if cb[b[j]] == 1}
        cands = [(i, pos_b[a[i]]) for i in range(a0, a1) if ca[a[i]] == 1 and a[i] in pos_b]
        if cands:
            keep = lis_indices([j for _, j in cands])
            pi, pj = a0, b0
            for k in keep:
                i, j = cands[k]
                pairs.append((i, j))
                stack.append((pi, i, pj, j))
                pi, pj = i + 1, j + 1
            stack.append((pi, a1, pj, b1))
        elif (a1 - a0) * (b1 - b0) <= 4_000_000:
            sm = difflib.SequenceMatcher(None, a[a0:a1], b[b0:b1], autojunk=False)
            for blk in sm.get_matching_blocks():
                for k in range(blk.size):
                    pairs.append((a0 + blk.a + k, b0 + blk.b + k))
    pairs.sort()
    return pairs


def order_score(ref_tokens, cand_tokens):
    """Share of unique n-gram anchors of the reference that appear in the same relative order in the candidate."""
    for n in (3, 2, 1):
        rg = [tuple(ref_tokens[i : i + n]) for i in range(len(ref_tokens) - n + 1)]
        cg = [tuple(cand_tokens[i : i + n]) for i in range(len(cand_tokens) - n + 1)]
        rc, cc = Counter(rg), Counter(cg)
        cpos = {g: i for i, g in enumerate(cg) if cc[g] == 1}
        seq = [cpos[g] for g in rg if rc[g] == 1 and g in cpos]
        if len(seq) >= 20 or n == 1:
            break
    if not seq:
        return {"score": None, "anchors": 0}
    return {"score": round(len(lis_indices(seq)) / len(seq), 4), "anchors": len(seq), "n": n}


# ── Fonts and colours ──────────────────────────────────────────────────────────────────────────────────────────────

FAMILY_ALIASES = {
    "helvetica": "arial",
    "arialmt": "arial",
    "liberationsans": "arial",
    "arimo": "arial",
    "helveticaneue": "arial",
    "times": "timesnewroman",
    "timesroman": "timesnewroman",
    "liberationserif": "timesnewroman",
    "tinos": "timesnewroman",
    "courier": "couriernew",
    "liberationmono": "couriernew",
    "cousine": "couriernew",
    "carlito": "calibri",
    "caladea": "cambria",
}
WEIGHT_FAMILIES = ("semibold", "semilight", "extrabold", "demibold", "condensed", "narrow", "medium", "light", "black", "thin")


def family_key(name):
    """Comparable family: lower-case, no spaces, PostScript suffixes removed, metric-compatible aliases merged."""
    if not name:
        return ""
    n = re.sub(r"^[A-Z]{6}\+", "", name)
    parts = re.split(r"[-,]", n, maxsplit=1)
    style = re.sub(r"[-_\s]", "", parts[1].lower()) if len(parts) == 2 else ""
    base = re.sub(r"\s+", "", parts[0]).lower()
    for suffix in ("psmt", "ps", "mt", "std", "lt", "pro"):
        if base.endswith(suffix) and len(base) > len(suffix) + 2:
            base = base[: -len(suffix)]
            break
    for tail in ("bolditalic", "boldoblique", "bold", "italic", "oblique", "regular"):
        if base.endswith(tail) and len(base) > len(tail) + 2:
            stem = base[: -len(tail)]
            if not (tail.startswith("bold") and stem.endswith(("semi", "demi", "extra", "ultra"))):  # "Segoe UI Semibold"
                base = stem
            break
    for weight in WEIGHT_FAMILIES:
        if weight in style.replace("italic", "").replace("oblique", ""):
            base += weight
            break
    return FAMILY_ALIASES.get(base, base)


def hex_rgb(h):
    h = (h or "000000").lstrip("#")
    try:
        return int(h[0:2], 16), int(h[2:4], 16), int(h[4:6], 16)
    except ValueError:
        return 0, 0, 0


def same_color(a, b, tol=40):
    ra, rb = hex_rgb(a), hex_rgb(b)
    return max(abs(x - y) for x, y in zip(ra, rb)) <= tol


def same_family(a, b):
    """Same family, allowing a weight written in the name on one side only ("Inter SemiBold" vs "Inter" + bold).
    Fonts without a usable name (Type3, unnamed) cannot be compared and count as agreeing."""
    ka, kb = family_key(a), family_key(b)
    if ka == kb or not ka or not kb or ka.startswith("type3") or kb.startswith("type3"):
        return True

    def base(k):
        # Only weights near regular/bold, which a converter may write as the base family (+ bold); Light/Thin are
        # visibly different faces and must keep their name.
        for weight in ("semibold", "demibold", "extrabold", "medium", "black"):
            if k.endswith(weight) and len(k) > len(weight) + 2:
                return k[: -len(weight)]
        return k

    return base(ka) == base(kb)


def same_style(a, b):
    """Per-attribute agreement of two word styles."""
    return {
        "family": same_family(a["family"], b["family"]),
        "size": abs(a["size"] - b["size"]) <= 0.6,
        "bold": a["bold"] is None or b["bold"] is None or a["bold"] == b["bold"],  # None: semibold, either way is fine
        "italic": a["italic"] is None or b["italic"] is None or a["italic"] == b["italic"],
        "color": a["color"] is None or b["color"] is None or same_color(a["color"], b["color"]),
    }


STYLE_KEYS = ("family", "size", "bold", "italic", "color")


def align_with_moves(a, b, passes=3):
    """align_tokens, then again on what is left unmatched on both sides, so blocks that moved are paired too."""
    pairs = align_tokens(a, b)
    for _ in range(passes):
        used_a = {i for i, _ in pairs}
        used_b = {j for _, j in pairs}
        rest_a = [i for i in range(len(a)) if i not in used_a]
        rest_b = [j for j in range(len(b)) if j not in used_b]
        if not rest_a or not rest_b:
            break
        more = align_tokens([a[i] for i in rest_a], [b[j] for j in rest_b])
        if not more:
            break
        pairs += [(rest_a[i], rest_b[j]) for i, j in more]
    return sorted(pairs)


def style_agreement(ref_words, cand_words, extra=()):
    """Pair two styled word lists (block moves allowed) and measure how often each style attribute agrees."""
    rt = [w["t"] for w in ref_words]
    ct = [w["t"] for w in cand_words]
    pairs = align_with_moves(rt, ct)
    counts = Counter()
    n = 0
    mismatches = Counter()
    examples = {}
    for i, j in pairs:
        a, b = ref_words[i], cand_words[j]
        if not re.search(r"\w", a["t"]):
            continue
        n += 1
        agree = same_style(a["style"], b["style"])
        for key in extra:
            agree[key] = a["style"].get(key) == b["style"].get(key)
        for key, ok in agree.items():
            if ok:
                counts[key] += 1
            else:
                label = f"{key}: {describe(a['style'], key)} -> {describe(b['style'], key)}"
                mismatches[label] += 1
                examples.setdefault(label, a["t"])
        if all(agree[k] for k in STYLE_KEYS):
            counts["all"] += 1
    if not n:
        return None
    out = {"words": n, "refWords": len(ref_words), "coverage": round(n / max(1, len([w for w in ref_words if re.search(r"\w", w["t"])])), 4)}
    for key in STYLE_KEYS + tuple(extra) + ("all",):
        out[key] = round(counts[key] / n, 4)
    out["topMismatches"] = [f"{k} ({v}, e.g. '{examples[k]}')" for k, v in mismatches.most_common(6)]
    return out


def describe(style, key):
    v = style.get(key)
    if key == "family":
        return v or "?"
    if key == "size":
        return f"{v:g}"
    return str(v)


# ── PDF side ───────────────────────────────────────────────────────────────────────────────────────────────────────


def pdf_font_style(span):
    font = span.get("font", "")
    flags = span.get("flags", 0)
    # MuPDF char flags (TEXT_COLLECT_STYLES): 1 strikeout, 2 underline, 8 bold, 16 filled, 32 stroked.
    # Filled + stroked glyphs are Word's synthetic bold (a face without a bold style, e.g. Calibri Light).
    cflags = span.get("char_flags", 0)
    low = font.lower()
    bold = bool(flags & 16) or bool(cflags & 8) or (cflags & 48) == 48 or bool(re.search(r"bold|black|heavy", low))
    if re.search(r"(semi|demi)[-_ ]?bold", low):
        bold = None  # a weight between regular and bold: Word may use the "Semibold" family or bold, both are right
    italic = bool(flags & 2) or bool(re.search(r"italic|oblique", low))
    # Stroked-only glyphs (outlined text): the span colour is the unused fill colour, the visible one is unknown.
    stroked_only = (cflags & 48) == 32
    # Invisible text (alpha 0: OCR layers, Chrome's selectable copy of outlined or Type3 text): only its size and
    # position mean something; its font, colour and weight are not what the reader sees.
    invisible = span.get("alpha", 255) == 0
    return {
        "family": "" if invisible else font,
        "size": round(span.get("size", 0) * 2) / 2,
        "bold": None if invisible else bold,
        "italic": None if invisible else italic,
        "color": None if stroked_only or invisible else "%06X" % (span.get("color", 0) & 0xFFFFFF),
        "underline": bool(cflags & 2),
        "strike": bool(cflags & 1),
    }


def pdf_analysis(path):
    """Text (per page), styled words, line and image counts of a PDF, in content order."""
    doc = fitz.open(path)
    pages_text, words, lines, images, sizes = [], [], 0, 0, []
    flags = fitz.TEXT_PRESERVE_WHITESPACE | fitz.TEXT_MEDIABOX_CLIP | getattr(fitz, "TEXT_COLLECT_STYLES", 0)
    for pno, page in enumerate(doc):
        sizes.append([round(page.rect.width, 1), round(page.rect.height, 1)])
        d = page.get_text("dict", flags=flags)
        chars = []  # (char, style)
        for block in d["blocks"]:
            if block.get("type") != 0:
                continue
            for line in block["lines"]:
                lines += 1
                for span in line["spans"]:
                    st = pdf_font_style(span)
                    for ch in span["text"]:
                        chars.append((ch, st))
                chars.append(("\n", None))
            chars.append(("\n", None))
        pages_text.append("".join(c for c, _ in chars))
        words.extend(styled_words(chars, page=pno))
        seen = set()
        for info in page.get_image_info(xrefs=True):
            bbox = fitz.Rect(info["bbox"]) & page.rect
            if bbox.width < 6 or bbox.height < 6:
                continue
            key = (info.get("xref"), tuple(round(v) for v in bbox))
            if key not in seen:
                seen.add(key)
                images += 1
    out = {"pages": doc.page_count, "text": "\n".join(pages_text), "words": words, "lines": lines, "images": images, "pageSizes": sizes}
    doc.close()
    return out


def styled_words(chars, page=None, para=None):
    """Split (char, style) pairs into words; a word's style is the majority style of its letters."""
    out, cur = [], []

    def flush():
        if not cur:
            return
        text = "".join(c for c, _ in cur)
        letters = [st for c, st in cur if st is not None and c.isalnum()] or [st for _, st in cur if st is not None]
        style = Counter(json.dumps(s, sort_keys=True) for s in letters).most_common(1)[0][0] if letters else None
        for tok in tokens_of(text):
            out.append({"t": tok, "style": json.loads(style) if style else None, "page": page})
        cur.clear()

    for ch, st in chars:
        if ch.isspace():
            flush()
        else:
            cur.append((ch, st))
    flush()
    return [w for w in out if w["style"] is not None]


# ── DOCX side ──────────────────────────────────────────────────────────────────────────────────────────────────────


def on_off(el):
    if el is None:
        return None
    return el.get(W_VAL) not in ("0", "false", "off", "none")


class DocxModel:
    """What Word sees in a .docx: paragraphs (body, tables, text boxes) with resolved run styles, tables, counts."""

    def __init__(self, path):
        self.path = path
        self.doc = Document(path)
        styles_el = self.doc.styles.element
        self.styles = {s.get(q("w:styleId")): s for s in styles_el.findall(q("w:style"))}
        self.default_style = {"paragraph": None, "character": None}
        for s in styles_el.findall(q("w:style")):
            if s.get(q("w:default")) in ("1", "true") and s.get(q("w:type")) in self.default_style:
                self.default_style[s.get(q("w:type"))] = s.get(q("w:styleId"))
        self.theme = self._theme_fonts()
        dd = styles_el.find(q("w:docDefaults"))
        self.default_rpr = dd.find(".//" + q("w:rPrDefault") + "/" + W_RPR) if dd is not None else None
        self.default_ppr = dd.find(".//" + q("w:pPrDefault") + "/" + W_PPR) if dd is not None else None
        self.numbering = self._numbering()
        self.paragraphs = []
        self.tables = []
        self.counts = Counter()
        self._style_cache = {}
        body = self.doc.element.body
        self._walk(body, "body", None)
        self._sections(body)
        self.header_paragraphs = []
        for sec in self.doc.sections:
            for part in (sec.header, sec.footer):
                try:
                    if part.is_linked_to_previous:
                        continue
                    self._walk(part._element, "header", None, out=self.header_paragraphs)
                except Exception:
                    pass

    # styles

    def _theme_fonts(self):
        fonts = {"major": "Calibri Light", "minor": "Calibri"}
        try:
            for rel in self.doc.part.rels.values():
                if rel.reltype.endswith("/theme"):
                    xml = rel.target_part.blob.decode("utf8", "replace")
                    for kind in ("major", "minor"):
                        m = re.search(r"<a:%sFont>\s*<a:latin typeface=\"([^\"]*)\"" % kind, xml)
                        if m:
                            fonts[kind] = m.group(1)
        except Exception:
            pass
        return fonts

    def _numbering(self):
        nums = {}
        try:
            el = self.doc.part.numbering_part.element
        except Exception:
            return nums
        abstract = {}
        for a in el.findall(q("w:abstractNum")):
            lvls = {}
            for lvl in a.findall(q("w:lvl")):
                fmt = lvl.find(q("w:numFmt"))
                lvls[int(lvl.get(q("w:ilvl")))] = fmt.get(W_VAL) if fmt is not None else "decimal"
            abstract[a.get(q("w:abstractNumId"))] = lvls
        for n in el.findall(q("w:num")):
            ref = n.find(q("w:abstractNumId"))
            if ref is not None:
                nums[n.get(q("w:numId"))] = abstract.get(ref.get(W_VAL), {})
        return nums

    def _chain(self, style_id):
        chain, seen = [], set()
        while style_id and style_id not in seen:
            seen.add(style_id)
            s = self.styles.get(style_id)
            if s is None:
                break
            chain.append(s)
            based = s.find(q("w:basedOn"))
            style_id = based.get(W_VAL) if based is not None else None
        return chain[::-1]

    def _rpr_props(self, rpr):
        props = {}
        if rpr is None:
            return props
        rf = rpr.find(q("w:rFonts"))
        if rf is not None:
            theme = rf.get(q("w:asciiTheme")) or rf.get(q("w:hAnsiTheme"))
            if theme:
                props["family"] = self.theme["major" if theme.startswith("major") else "minor"]
            elif rf.get(q("w:ascii")) or rf.get(q("w:hAnsi")):
                props["family"] = rf.get(q("w:ascii")) or rf.get(q("w:hAnsi"))
        sz = rpr.find(q("w:sz"))
        if sz is not None and sz.get(W_VAL):
            try:
                props["size"] = int(float(sz.get(W_VAL))) / 2
            except ValueError:
                pass
        for key, tag in (("bold", "w:b"), ("italic", "w:i"), ("strike", "w:strike"), ("caps", "w:caps"), ("smallCaps", "w:smallCaps")):
            v = on_off(rpr.find(q(tag)))
            if v is not None:
                props[key] = v
        u = rpr.find(q("w:u"))
        if u is not None:
            props["underline"] = u.get(W_VAL, "single") != "none"
        c = rpr.find(q("w:color"))
        if c is not None and c.get(W_VAL):
            val = c.get(W_VAL)
            props["color"] = "000000" if val.lower() == "auto" else val.upper()
        va = rpr.find(q("w:vertAlign"))
        if va is not None:
            props["script"] = {"superscript": "sup", "subscript": "sub"}.get(va.get(W_VAL), "")
        return props

    def _style_rpr(self, style_id):
        key = ("r", style_id)
        if key not in self._style_cache:
            props = {}
            for s in self._chain(style_id):
                props.update(self._rpr_props(s.find(W_RPR)))
            self._style_cache[key] = props
        return self._style_cache[key]

    def _style_ppr(self, style_id):
        key = ("p", style_id)
        if key not in self._style_cache:
            props = {}
            for s in self._chain(style_id):
                props.update(self._ppr_props(s.find(W_PPR), s))
            self._style_cache[key] = props
        return self._style_cache[key]

    def _ppr_props(self, ppr, style=None):
        props = {}
        if style is not None:
            name = style.find(q("w:name"))
            m = re.match(r"heading\s*(\d)$", (name.get(W_VAL) if name is not None else "") or "", re.I)
            if m:
                props["heading"] = int(m.group(1))
        if ppr is None:
            return props
        jc = ppr.find(q("w:jc"))
        if jc is not None:
            props["align"] = {"both": "justify", "distribute": "justify", "start": "left", "end": "right", "center": "center", "right": "right", "left": "left"}.get(
                jc.get(W_VAL), "left"
            )
        ol = ppr.find(q("w:outlineLvl"))
        if ol is not None:
            lvl = int(ol.get(W_VAL, "9"))
            props["heading"] = lvl + 1 if lvl < 9 else 0
        num = ppr.find(q("w:numPr"))
        if num is not None:
            nid = num.find(q("w:numId"))
            ilvl = num.find(q("w:ilvl"))
            if nid is not None:
                props["numId"] = nid.get(W_VAL)
            if ilvl is not None:
                props["ilvl"] = int(ilvl.get(W_VAL, "0"))
        if ppr.find(q("w:framePr")) is not None:
            props["frame"] = True
        return props

    def run_style(self, r, para_style):
        props = {"family": "Times New Roman", "size": 10.0, "bold": False, "italic": False, "color": "000000", "underline": False, "script": ""}
        props.update(self._rpr_props(self.default_rpr))
        props.update(self._style_rpr(para_style))
        rpr = r.find(W_RPR)
        rstyle = rpr.find(q("w:rStyle")) if rpr is not None else None
        char_style = rstyle.get(W_VAL) if rstyle is not None else self.default_style["character"]
        if char_style:
            props.update(self._style_rpr(char_style))
        props.update(self._rpr_props(rpr))
        return {k: props.get(k) for k in ("family", "size", "bold", "italic", "color", "underline", "script")}

    # document walk

    def _walk(self, container, where, cell, out=None):
        out = self.paragraphs if out is None else out
        for child in container:
            tag = child.tag
            if tag == W_P:
                self._paragraph(child, where, cell, out)
            elif tag == W_TBL:
                self._table(child, where, out)
            elif tag == W_SDT:
                content = child.find(W_SDTCONTENT)
                if content is not None:
                    self._walk(content, where, cell, out)
            elif tag == MC_ALT:
                choice = child.find(MC_CHOICE)
                if choice is not None:
                    self._walk(choice, where, cell, out)
            elif tag == q("w:customXml"):
                self._walk(child, where, cell, out)
            elif tag == q("w:sectPr"):
                pass

    def _paragraph(self, p, where, cell, out):
        ppr = p.find(W_PPR)
        pstyle_el = ppr.find(q("w:pStyle")) if ppr is not None else None
        pstyle = pstyle_el.get(W_VAL) if pstyle_el is not None else self.default_style["paragraph"]
        props = {"align": "left", "heading": 0}
        props.update(self._ppr_props(self.default_ppr))
        props.update(self._style_ppr(pstyle))
        props.update(self._ppr_props(ppr))
        if ppr is not None and ppr.find(q("w:sectPr")) is not None:
            self.counts["sectionBreaks"] += 1
        lst = None
        if props.get("numId") and props["numId"] != "0":
            fmt = self.numbering.get(props["numId"], {}).get(props.get("ilvl", 0), "decimal")
            if fmt != "none":
                lst = {"kind": "bullet" if fmt == "bullet" else "number", "level": props.get("ilvl", 0)}
        chars = []  # (char, style) ; "\n" = soft line break
        nested = []

        def visit(el, style):
            for child in el:
                tag = child.tag
                if tag in SKIP_TEXT:
                    continue
                if tag == W_R:
                    visit(child, self.run_style(child, pstyle))
                elif tag == W_T:
                    for ch in child.text or "":
                        chars.append((ch, style))
                elif tag == W_TAB or tag == q("w:ptab"):
                    chars.append(("\t", style))
                    self.counts["tabs"] += 1
                elif tag == W_BR:
                    kind = child.get(q("w:type"))
                    if kind in (None, "textWrapping"):
                        chars.append(("\n", style))
                        self.counts["lineBreaks"] += 1
                    else:
                        chars.append((" ", style))
                        self.counts["pageBreaks" if kind == "page" else "columnBreaks"] += 1
                elif tag == W_CR:
                    chars.append(("\n", style))
                    self.counts["lineBreaks"] += 1
                elif tag == q("w:noBreakHyphen"):
                    chars.append(("-", style))
                elif tag == W_SYM:
                    try:
                        chars.append((chr(int(child.get(q("w:char")), 16)), style))
                    except (TypeError, ValueError):
                        pass
                elif tag == W_TXBX:
                    nested.append(child)
                    self.counts["textboxes"] += 1
                elif tag == MC_ALT:
                    choice = child.find(MC_CHOICE)
                    if choice is not None:
                        visit(choice, style)
                elif tag == MC_FALLBACK:
                    continue
                elif tag == q("w:hyperlink"):
                    self.counts["hyperlinks"] += 1
                    visit(child, style)
                elif tag == q("w:drawing") or tag == q("w:pict") or tag == q("w:object"):
                    self._drawing(child)
                    visit(child, style)
                elif tag == q("w:fldSimple"):
                    if "HYPERLINK" in (child.get(q("w:instr")) or ""):
                        self.counts["hyperlinks"] += 1
                    visit(child, style)
                else:
                    visit(child, style)

        visit(p, None)
        text = "".join(c for c, _ in chars)
        if "HYPERLINK" in "".join(t.text or "" for t in p.iter(q("w:instrText"))):
            self.counts["hyperlinks"] += 1
        words = []
        soft = False
        cur = []

        def flush():
            nonlocal soft
            if cur:
                for k, w in enumerate(styled_words(cur)):
                    w["soft"] = soft and k == 0
                    words.append(w)
                cur.clear()
                soft = False

        for ch, st in chars:
            if ch == "\n":
                flush()
                soft = True
            elif ch.isspace():
                flush()
            else:
                cur.append((ch, st))
        flush()
        para = {
            "index": len(out),
            "where": where,
            "cell": cell,
            "text": text,
            "words": words,
            "align": props.get("align", "left"),
            "heading": props.get("heading", 0) if props.get("heading", 0) <= 9 else 0,
            "list": lst,
            "frame": bool(props.get("frame")),
        }
        if para["frame"]:
            self.counts["frames"] += 1
        out.append(para)
        for box in nested:
            self._walk(box, "textbox", None, out)

    def _drawing(self, el):
        """Counts the pictures of one drawing (text boxes inside it are counted when their paragraphs are walked)."""

        def pruned(node, tag):
            for child in node:
                if child.tag in (W_TXBX, MC_FALLBACK):
                    continue
                if child.tag == tag:
                    yield child
                yield from pruned(child, tag)

        def has_picture(node):
            return any(True for _ in pruned(node, q("pic:pic")))

        def extent(node):
            ext = node.find(q("wp:extent"))
            if ext is None:
                return None
            return int(ext.get("cx", 0)) / 12700, int(ext.get("cy", 0)) / 12700

        for anchor in pruned(el, q("wp:anchor")):
            if has_picture(anchor):
                self.counts["anchoredImages"] += 1
                size = extent(anchor)
                behind = anchor.get("behindDoc") in ("1", "true")
                if behind:
                    self.counts["behindImages"] += 1
                if size and behind and size[0] > 400 and size[1] > 400:
                    self.counts["pageBackgrounds"] += 1
                elif size and min(size) < 3:
                    self.counts["rulePictures"] += 1  # a line drawn as a picture (underline, rule, border)
        for inline in pruned(el, q("wp:inline")):
            if has_picture(inline):
                self.counts["inlineImages"] += 1
                size = extent(inline)
                if size and min(size) < 3:
                    self.counts["rulePictures"] += 1
        for _ in pruned(el, q("v:imagedata")):
            self.counts["vmlImages"] += 1

    @staticmethod
    def _visible_edges(borders):
        """None when the element says nothing, else whether any edge is drawn."""
        if borders is None:
            return None
        return any(edge.get(W_VAL) not in ("nil", "none", None) for edge in borders)

    def _table_borders(self, tbl):
        """Whether the table draws any border (direct formatting, its table style, or cell borders)."""
        tbl_pr = tbl.find(q("w:tblPr"))
        direct = self._visible_edges(tbl_pr.find(q("w:tblBorders"))) if tbl_pr is not None else None
        if direct is not None:
            if direct:
                return True
        else:
            style = tbl_pr.find(q("w:tblStyle")) if tbl_pr is not None else None
            for s in self._chain(style.get(W_VAL) if style is not None else None):
                s_pr = s.find(q("w:tblPr"))
                v = self._visible_edges(s_pr.find(q("w:tblBorders"))) if s_pr is not None else None
                if v:
                    return True
        return any(self._visible_edges(b) for b in tbl.iter(q("w:tcBorders")))

    def _table(self, tbl, where, out):
        table = {"index": len(self.tables), "rows": [], "nested": where == "table", "borders": self._table_borders(tbl)}
        self.tables.append(table)
        for r, tr in enumerate(tbl.findall(W_TR)):
            row = []
            tr_pr = tr.find(q("w:trPr"))
            before = tr_pr.find(q("w:gridBefore")) if tr_pr is not None else None
            if before is not None and int(before.get(W_VAL, "0") or 0) > 0:
                row.append({"text": "", "span": int(before.get(W_VAL)), "vmerge": None, "shade": None, "filler": True})
            cells = [c for c in tr if c.tag == W_TC]
            for sdt in tr.findall(W_SDT):
                content = sdt.find(W_SDTCONTENT)
                if content is not None:
                    cells.extend(c for c in content if c.tag == W_TC)
            for c, tc in enumerate(cells):
                tcpr = tc.find(q("w:tcPr"))
                span, vmerge, shade = 1, None, None
                if tcpr is not None:
                    gs = tcpr.find(q("w:gridSpan"))
                    if gs is not None:
                        span = int(gs.get(W_VAL, "1"))
                    vm = tcpr.find(q("w:vMerge"))
                    if vm is not None:
                        vmerge = vm.get(W_VAL, "continue")
                    shd = tcpr.find(q("w:shd"))
                    if shd is not None and shd.get(q("w:fill")) not in (None, "auto", "FFFFFF"):
                        shade = shd.get(q("w:fill"))
                start = len(out)
                self._walk(tc, "table", (table["index"], r, c), out)
                text = " ".join(p["text"] for p in out[start:] if p["where"] == "table")
                row.append({"text": norm_text(text), "span": span, "vmerge": vmerge, "shade": shade})
            after = tr_pr.find(q("w:gridAfter")) if tr_pr is not None else None
            if after is not None and int(after.get(W_VAL, "0") or 0) > 0:
                row.append({"text": "", "span": int(after.get(W_VAL)), "vmerge": None, "shade": None, "filler": True})
            table["rows"].append(row)
        grid = []
        for r, row in enumerate(table["rows"]):
            line = []
            for cell in row:
                text = cell["text"]
                if cell["vmerge"] == "continue" and r > 0:
                    col = len(line)
                    above = grid[r - 1][col] if col < len(grid[r - 1]) else ""
                    text = above
                for _ in range(cell["span"]):
                    line.append(text)
            grid.append(line)
        table["grid"] = grid
        table["cols"] = max((len(r) for r in grid), default=0)
        table["merged"] = sum(1 for row in table["rows"] for c in row if not c.get("filler") and (c["span"] > 1 or c["vmerge"]))
        table["shaded"] = sum(1 for row in table["rows"] for c in row if c["shade"])
        # A table with neither borders nor shading is invisible on the page (layout only).
        table["visible"] = bool(table["borders"] or table["shaded"])

    def _sections(self, body):
        sects = list(body.iter(q("w:sectPr")))
        self.counts["sections"] = len(sects)
        cols = 1
        landscape = 0
        for s in sects:
            c = s.find(q("w:cols"))
            if c is not None:
                cols = max(cols, int(c.get(q("w:num"), "1") or 1))
            sz = s.find(q("w:pgSz"))
            if sz is not None and sz.get(q("w:w")) and sz.get(q("w:h")) and int(sz.get(q("w:w"))) > int(sz.get(q("w:h"))):
                landscape += 1
        self.counts["maxColumns"] = cols
        self.counts["landscapeSections"] = landscape

    # summaries

    def body_words(self):
        """Styled words of body, table and text-box paragraphs, with their paragraph index."""
        out = []
        for p in self.paragraphs:
            for w in p["words"]:
                out.append(dict(w, para=p["index"]))
        return out

    def text(self, headers=False):
        paras = self.paragraphs + (self.header_paragraphs if headers else [])
        return "\n".join(p["text"] for p in paras)

    def structure(self):
        nonempty = [p for p in self.paragraphs if p["text"].strip()]
        chars = Counter()
        for p in nonempty:
            # Framed paragraphs (w:framePr) are positioned on the page: not flowing text, like text boxes.
            chars["frame" if p["frame"] else p["where"]] += len(p["text"].strip())
        total = sum(chars.values())
        top_tables = [t for t in self.tables if not t["nested"]]
        return {
            "paragraphs": len(nonempty),
            "bodyParagraphs": sum(1 for p in nonempty if p["where"] == "body"),
            "tableParagraphs": sum(1 for p in nonempty if p["where"] == "table"),
            "textboxParagraphs": sum(1 for p in nonempty if p["where"] == "textbox"),
            "emptyParagraphs": len(self.paragraphs) - len(nonempty),
            "tables": len(top_tables),
            "layoutTables": sum(1 for t in top_tables if not t["visible"]),
            "nestedTables": len(self.tables) - len(top_tables),
            "tableDims": [dims_label(t) for t in top_tables],
            "mergedCells": sum(t["merged"] for t in self.tables),
            "images": self.counts["inlineImages"] + self.counts["anchoredImages"] + self.counts["vmlImages"] - self.counts["pageBackgrounds"] - self.counts["rulePictures"],
            "rulePictures": self.counts["rulePictures"],
            "inlineImages": self.counts["inlineImages"],
            "anchoredImages": self.counts["anchoredImages"],
            "pageBackgrounds": self.counts["pageBackgrounds"],
            "textboxes": self.counts["textboxes"],
            "frames": self.counts["frames"],
            "sections": self.counts["sections"],
            "maxColumns": self.counts["maxColumns"],
            "landscapeSections": self.counts["landscapeSections"],
            "headings": sum(1 for p in nonempty if p["heading"]),
            "listItems": sum(1 for p in nonempty if p["list"]),
            "hyperlinks": self.counts["hyperlinks"],
            "lineBreaks": self.counts["lineBreaks"],
            "tabs": self.counts["tabs"],
            "headerParagraphs": sum(1 for p in self.header_paragraphs if p["text"].strip()),
            "flowRatio": round((chars["body"] + chars["table"]) / total, 4) if total else None,
        }


def dims_label(table):
    """rows x columns; '*' marks an invisible (layout) table."""
    return f"{len(table['grid'])}x{table['cols']}{'' if table['visible'] else '*'}"


# ── Ground-truth comparisons ───────────────────────────────────────────────────────────────────────────────────────


def paragraph_metrics(gt, cand):
    """Paragraph breaks (precision/recall/F1) and per-paragraph alignment/heading/list agreement vs the ground truth."""
    gw, cw = gt.body_words(), cand.body_words()
    pairs = align_tokens([w["t"] for w in gw], [w["t"] for w in cw])
    if not pairs:
        return None
    tp = fp = fn = tp_strict = fp_strict = fn_strict = 0
    for (i1, j1), (i2, j2) in zip(pairs, pairs[1:]):
        if i2 != i1 + 1:
            continue
        g_break = gw[i1]["para"] != gw[i2]["para"]
        c_para = cw[j1]["para"] != cw[j2]["para"]
        c_soft = any(cw[k]["soft"] for k in range(j1 + 1, j2 + 1))
        c_break = c_para or c_soft
        tp += g_break and c_break
        fp += (not g_break) and c_break
        fn += g_break and not c_break
        tp_strict += g_break and c_para
        fp_strict += (not g_break) and c_para
        fn_strict += g_break and not c_para

    def f1(t, p, n):
        prec = t / (t + p) if t + p else 1.0
        rec = t / (t + n) if t + n else 1.0
        return round(prec, 4), round(rec, 4), round(2 * prec * rec / (prec + rec), 4) if prec + rec else 0.0

    prec, rec, f = f1(tp, fp, fn)
    _, _, f_strict = f1(tp_strict, fp_strict, fn_strict)
    # map each ground-truth paragraph to the candidate paragraph holding most of its words
    votes = {}
    for i, j in pairs:
        votes.setdefault(gw[i]["para"], Counter())[cw[j]["para"]] += 1
    cpar = {p["index"]: p for p in cand.paragraphs}
    align_ok = align_n = head_ok = head_n = list_ok = list_n = 0
    for gp in gt.paragraphs:
        if gp["index"] not in votes or not gp["words"]:
            continue
        cp = cpar[votes[gp["index"]].most_common(1)[0][0]]
        ga, ca = gp["align"], cp["align"]
        if len(gp["words"]) < 14 and {ga, ca} <= {"left", "justify"}:
            ca = ga  # a one-line justified paragraph looks left-aligned
        align_n += 1
        align_ok += ga == ca
        head_n += 1
        head_ok += (gp["heading"] or 0) == (cp["heading"] or 0)
        if gp["list"] or cp["list"]:
            list_n += 1
            list_ok += bool(gp["list"] and cp["list"] and gp["list"]["kind"] == cp["list"]["kind"] and gp["list"]["level"] == cp["list"]["level"])
    return {
        "breakPrecision": prec,
        "breakRecall": rec,
        "breakF1": f,
        "breakF1Strict": f_strict,
        "alignment": round(align_ok / align_n, 4) if align_n else None,
        "heading": round(head_ok / head_n, 4) if head_n else None,
        "list": round(list_ok / list_n, 4) if list_n else None,
        "gtParagraphs": sum(1 for p in gt.paragraphs if p["text"].strip()),
        "paragraphs": sum(1 for p in cand.paragraphs if p["text"].strip()),
    }


def table_metrics(gt, cand):
    """Visible ground-truth tables (borders or shading) matched to the candidate's tables by cell text."""
    top = [t for t in gt.tables if not t["nested"]]
    gts = [t for t in top if t["visible"]]
    layout = len(top) - len(gts)
    cts = [t for t in cand.tables if not t["nested"]]
    if not gts:
        return {"gt": [], "layoutTables": layout, "ours": [dims_label(t) for t in cts], "matched": 0, "cellsPos": None, "cellsBag": None, "dims": None}
    used = set()
    results = []
    for g in gts:
        gcells = [c for row in g["grid"] for c in row if c]
        best, best_score = None, 0.0
        for k, c in enumerate(cts):
            if k in used:
                continue
            ccells = [x for row in c["grid"] for x in row if x]
            inter = sum((Counter(gcells) & Counter(ccells)).values())
            score = inter / max(1, len(set(gcells) | set(ccells)))
            if score > best_score:
                best, best_score = k, score
        if best is None or best_score < 0.2:
            results.append({"gt": f"{len(g['grid'])}x{g['cols']}", "ours": None, "pos": 0.0, "bag": 0.0, "dims": False})
            continue
        used.add(best)
        c = cts[best]
        positions = [(r, k) for r, row in enumerate(g["grid"]) for k, _ in enumerate(row)]
        pos_ok = sum(1 for r, k in positions if r < len(c["grid"]) and k < len(c["grid"][r]) and c["grid"][r][k] == g["grid"][r][k])
        ccells = Counter(x for row in c["grid"] for x in row if x)
        bag_ok = sum((Counter(gcells) & ccells).values())
        results.append(
            {
                "gt": f"{len(g['grid'])}x{g['cols']}",
                "ours": f"{len(c['grid'])}x{c['cols']}",
                "pos": round(pos_ok / max(1, len(positions)), 4),
                "bag": round(bag_ok / max(1, len(gcells)), 4),
                "dims": len(g["grid"]) == len(c["grid"]) and g["cols"] == c["cols"],
                "merged": [g["merged"], c["merged"]],
            }
        )
    return {
        "gt": [r["gt"] for r in results],
        "layoutTables": layout,
        "ours": [dims_label(t) for t in cts],
        "matched": sum(1 for r in results if r["ours"]),
        "cellsPos": round(sum(r["pos"] for r in results) / len(results), 4),
        "cellsBag": round(sum(r["bag"] for r in results) / len(results), 4),
        "dims": round(sum(1 for r in results if r["dims"]) / len(results), 4),
        "detail": results,
    }


def docx_style_words(model):
    return [dict(w, style={k: w["style"][k] for k in ("family", "size", "bold", "italic", "color", "underline", "script")}) for w in model.body_words()]


SCRIPT_SCALE = 0.63  # Word draws superscript/subscript runs at ~63 % of their font size (measured: 11 pt -> 6.96 pt)


def docx_words_as_drawn(model):
    """The .docx words (body, tables, text boxes, headers) with the size Word draws them at, for comparing with PDF text."""
    out = []
    for w in model.body_words() + [dict(w, para=-1) for p in model.header_paragraphs for w in p["words"]]:
        st = dict(w["style"])
        if st.get("script"):
            st["size"] = round(st["size"] * SCRIPT_SCALE * 2) / 2
        out.append(dict(w, style=st))
    return out


# ── Rendering ──────────────────────────────────────────────────────────────────────────────────────────────────────


def render_gray(page, width_px, height_px):
    m = fitz.Matrix(width_px / page.rect.width, height_px / page.rect.height)
    pix = page.get_pixmap(matrix=m, colorspace=fitz.csGRAY, alpha=False)
    arr = np.frombuffer(pix.samples, dtype=np.uint8).reshape(pix.height, pix.stride)[:, : pix.width].astype(np.float32)
    if arr.shape != (height_px, width_px):
        out = np.full((height_px, width_px), 255, np.float32)
        h, w = min(height_px, arr.shape[0]), min(width_px, arr.shape[1])
        out[:h, :w] = arr[:h, :w]
        arr = out
    return arr


def box_blur(a, r=1):
    p = np.pad(a, r, mode="edge")
    c = np.cumsum(np.cumsum(p, axis=0), axis=1)
    c = np.pad(c, ((1, 0), (1, 0)))
    k = 2 * r + 1
    s = c[k:, k:] - c[:-k, k:] - c[k:, :-k] + c[:-k, :-k]
    return s / (k * k)


def dilate(mask, r=2):
    out = mask.copy()
    h, w = mask.shape
    for dy in range(-r, r + 1):
        for dx in range(-r, r + 1):
            if dx == 0 and dy == 0:
                continue
            ys, yd = (slice(dy, h), slice(0, h - dy)) if dy >= 0 else (slice(0, h + dy), slice(-dy, h))
            xs, xd = (slice(dx, w), slice(0, w - dx)) if dx >= 0 else (slice(0, w + dx), slice(-dx, w))
            out[yd, xd] |= mask[ys, xs]
    return out


def shifted(arr, dy, dx, fill):
    """arr moved by (dy, dx) pixels, uncovered area filled with `fill`."""
    out = np.full_like(arr, fill)
    h, w = arr.shape
    ys, yd = (slice(0, h - dy), slice(dy, h)) if dy >= 0 else (slice(-dy, h), slice(0, h + dy))
    xs, xd = (slice(0, w - dx), slice(dx, w)) if dx >= 0 else (slice(-dx, w), slice(0, w + dx))
    out[yd, xd] = arr[ys, xs]
    return out


def best_shift(ia, ib, max_shift):
    """(dy, dx) that moves ink mask `ib` onto `ia` with the most overlap (FFT cross-correlation, |d| <= max_shift)."""
    if not ia.any() or not ib.any():
        return 0, 0
    h, w = ia.shape
    ph, pw = h + max_shift, w + max_shift
    fa = np.fft.rfft2(box_blur(ia.astype(np.float32)), s=(ph, pw))
    fb = np.fft.rfft2(box_blur(ib.astype(np.float32)), s=(ph, pw))
    corr = np.fft.irfft2(fa * np.conj(fb), s=(ph, pw))
    best, arg = -1.0, (0, 0)
    for dy in range(-max_shift, max_shift + 1):
        row = corr[dy % ph]
        for dx in range(-max_shift, max_shift + 1):
            v = row[dx % pw] - 1e-6 * (dy * dy + dx * dx)  # prefer the smallest shift among equals
            if v > best:
                best, arg = v, (dy, dx)
    return arg


def ink_f1(ia, ib, r):
    na, nb = int(ia.sum()), int(ib.sum())
    if na == 0 and nb == 0:
        return 1.0
    if na == 0 or nb == 0:
        return 0.0
    prec = float((ib & dilate(ia, r)).sum()) / nb
    rec = float((ia & dilate(ib, r)).sum()) / na
    return 2 * prec * rec / (prec + rec) if prec + rec else 0.0


def visual_metrics(orig_path, cand_path, dpi=60, max_shift_pt=18):
    """Per page, after the best global shift (reported in points): mean absolute difference of blurred grey renders
    (0 = identical) and ink F1 within ~2.4 pt (inkF1) or ~7 pt (inkF1Loose); inkF1Raw is without the shift."""
    o, c = fitz.open(orig_path), fitz.open(cand_path)
    pages = []
    max_shift = max(1, round(max_shift_pt * dpi / 72))
    for i in range(max(o.page_count, c.page_count)):
        if i >= o.page_count or i >= c.page_count:
            pages.append({"page": i + 1, "diff": 1.0, "inkF1": 0.0, "inkF1Loose": 0.0, "inkF1Raw": 0.0, "missing": "orig" if i >= o.page_count else "ours"})
            continue
        po = o[i]
        w, h = max(1, round(po.rect.width * dpi / 72)), max(1, round(po.rect.height * dpi / 72))
        a, b = render_gray(po, w, h), render_gray(c[i], w, h)
        ia, ib = a < 170, b < 170
        raw = ink_f1(ia, ib, 2)
        dy, dx = best_shift(ia, ib, max_shift)
        if dy or dx:
            b = shifted(b, dy, dx, 255.0)
            ib = b < 170
        diff = float(np.mean(np.abs(box_blur(a) - box_blur(b)))) / 255
        pages.append(
            {
                "page": i + 1,
                "diff": round(diff, 4),
                "inkF1": round(ink_f1(ia, ib, 2), 4),
                "inkF1Loose": round(ink_f1(ia, ib, 6), 4),
                "inkF1Raw": round(raw, 4),
                "shiftPt": [round(dx * 72 / dpi, 1), round(dy * 72 / dpi, 1)],
            }
        )
    o.close()
    c.close()
    if not pages:
        return None
    avg = lambda key: round(sum(p[key] for p in pages) / len(pages), 4)  # noqa: E731
    return {"diff": avg("diff"), "inkF1": avg("inkF1"), "inkF1Loose": avg("inkF1Loose"), "inkF1Raw": avg("inkF1Raw"), "pages": pages}


def load_font(size):
    for name in ("arial.ttf", "C:/Windows/Fonts/arial.ttf", "DejaVuSans.ttf"):
        try:
            return ImageFont.truetype(name, size)
        except OSError:
            continue
    return ImageFont.load_default()


def render_rgb(doc, index, dpi):
    if doc is None or index >= doc.page_count:
        return None
    pix = doc[index].get_pixmap(matrix=fitz.Matrix(dpi / 72, dpi / 72), alpha=False)
    return Image.frombytes("RGB", (pix.width, pix.height), pix.samples)


def overlay(a, b):
    """Original in red, candidate in blue: shifted text shows as coloured fringes, matching ink as dark purple."""
    if a is None or b is None:
        return None
    b = b.resize(a.size)
    ga = np.asarray(a.convert("L"), dtype=np.uint8)
    gb = np.asarray(b.convert("L"), dtype=np.uint8)
    rgb = np.stack([gb, np.minimum(ga, gb), ga], axis=-1)
    return Image.fromarray(rgb, "RGB")


def build_sheets(out_dir, orig_path, panels, dpi=70, max_pages=12, with_overlay=False, notes=None):
    """One PNG per page: original | each candidate rendered by Word (| overlay original/ours)."""
    os.makedirs(os.path.join(out_dir, "sheets"), exist_ok=True)
    docs = [("Original (PDF)", fitz.open(orig_path))] + [(label, fitz.open(p) if p and os.path.exists(p) else None) for label, p in panels]
    n = max(d.page_count for _, d in docs if d is not None)
    font, small = load_font(15), load_font(12)
    files = []
    for i in range(min(n, max_pages)):
        imgs = [(label, render_rgb(d, i, dpi)) for label, d in docs]
        if with_overlay and len(imgs) > 1:
            imgs.append(("Superposición (rojo = original, azul = nuestro)", overlay(imgs[0][1], imgs[1][1])))
        ref = next((im for _, im in imgs if im is not None), None)
        pw = max((im.width for _, im in imgs if im is not None), default=600)
        ph = max((im.height for _, im in imgs if im is not None), default=800)
        gap, top = 14, 44
        sheet = Image.new("RGB", (len(imgs) * (pw + gap) + gap, ph + top + gap), (228, 230, 235))
        draw = ImageDraw.Draw(sheet)
        for k, (label, im) in enumerate(imgs):
            x = gap + k * (pw + gap)
            draw.text((x, 6), label, fill=(20, 20, 20), font=font)
            note = (notes or {}).get(label, {}).get(i + 1)
            if note:
                draw.text((x, 25), note, fill=(90, 90, 90), font=small)
            if im is None:
                draw.rectangle([x, top, x + pw, top + ph], fill=(245, 245, 245), outline=(200, 80, 80))
                draw.text((x + 20, top + 20), "sin página", fill=(200, 80, 80), font=font)
            else:
                sheet.paste(im, (x, top))
                draw.rectangle([x - 1, top - 1, x + im.width, top + im.height], outline=(150, 150, 150))
        name = f"sheets/p{i + 1}.png"
        sheet.save(os.path.join(out_dir, name), optimize=True)
        files.append(name)
        del ref
    for _, d in docs:
        if d is not None:
            d.close()
    return files


# ── One document ───────────────────────────────────────────────────────────────────────────────────────────────────


def score_candidate(cand, orig, gt_model, visual_dpi):
    res = {}
    model = None
    if cand.get("docx") and os.path.exists(cand["docx"]):
        try:
            model = DocxModel(cand["docx"])
            res["structure"] = model.structure()
        except Exception as err:
            res["docxError"] = f"{type(err).__name__}: {err}"
    rendered = None
    if cand.get("pdf") and os.path.exists(cand["pdf"]):
        rendered = pdf_analysis(cand["pdf"])
        res["pages"] = rendered["pages"]
        res["pageSizes"] = rendered["pageSizes"][:3]
    ref_tokens = tokens_of(orig["text"])
    if rendered is not None:
        cand_tokens = tokens_of(rendered["text"])
        res["text"] = {"sim": text_similarity(orig["text"], rendered["text"]), "words": bag_f1(ref_tokens, cand_tokens), "source": "word-render"}
        res["order"] = order_score(ref_tokens, cand_tokens)
        # No underline here: MuPDF reads table rules under text as underlines.
        res["styleRendered"] = style_agreement(orig["words"], rendered["words"])
        try:
            res["visual"] = visual_metrics(orig["path"], cand["pdf"], visual_dpi)
        except Exception as err:
            res["visual"] = {"error": str(err)}
    if model is not None:
        docx_text = model.text(headers=True)
        res.setdefault("text", {"source": "docx"})
        res["text"]["simDocx"] = text_similarity(orig["text"], docx_text)
        if rendered is None:
            res["text"]["sim"] = res["text"]["simDocx"]
            res["text"]["words"] = bag_f1(ref_tokens, tokens_of(docx_text))
            res["order"] = order_score(ref_tokens, tokens_of(docx_text))
        # What the .docx asks for, against the original PDF: independent of the fonts installed where Word renders.
        res["styleDocx"] = style_agreement(orig["words"], docx_words_as_drawn(model))
        if gt_model is not None:
            res["style"] = style_agreement(docx_style_words(gt_model), docx_style_words(model), extra=("underline", "script"))
            res["paragraphs"] = paragraph_metrics(gt_model, model)
            res["tables"] = table_metrics(gt_model, model)
    res["score"] = composite(res)
    return res


COMPOSITE = (
    ("text", "sim"),
    ("order", "score"),
    ("styleDocx", "all"),  # what the .docx asks for: independent of the fonts installed where Word renders
    ("style", "all"),
    ("visual", "inkF1"),
    ("paragraphs", "breakF1"),
    ("tables", "cellsPos"),
)


def composite(res):
    """Mean of the headline metrics that exist (0..1, higher is better)."""
    parts = [v for v in ((res.get(group) or {}).get(key) for group, key in COMPOSITE) if isinstance(v, (int, float))]
    return round(sum(parts) / len(parts), 4) if parts else None


def run_job(job):
    t0 = time.time()
    out_dir = job["outDir"]
    os.makedirs(out_dir, exist_ok=True)
    orig = pdf_analysis(job["orig"])
    orig["path"] = job["orig"]
    result = {
        "id": job["id"],
        "orig": {
            "pages": orig["pages"],
            "chars": len(norm_text(orig["text"])),
            "words": len(orig["words"]),
            "lines": orig["lines"],
            "images": orig["images"],
            "pageSizes": orig["pageSizes"][:3],
        },
    }
    gt_model = None
    if job.get("gt") and os.path.exists(job["gt"]):
        gt_model = DocxModel(job["gt"])
        result["gt"] = gt_model.structure()
    for name, cand in (job.get("candidates") or {}).items():
        if not cand:
            continue
        try:
            result[name] = score_candidate(cand, orig, gt_model, job.get("visualDpi", 60))
        except Exception as err:
            import traceback

            result[name] = {"error": f"{type(err).__name__}: {err}", "trace": traceback.format_exc()[-1500:]}
    sheets = job.get("sheets")
    if sheets:
        panels = []
        labels = {"ours": "LocalPDF (.docx en Word)", "word": "Word (conversión propia)"}
        notes = {}
        for name, cand in (job.get("candidates") or {}).items():
            if cand and cand.get("pdf"):
                panels.append((labels.get(name, name), cand["pdf"]))
                vis = (result.get(name) or {}).get("visual") or {}
                notes[labels.get(name, name)] = {
                    p["page"]: f"inkF1 {p['inkF1']:.2f} (sin alinear {p['inkF1Raw']:.2f}) · desplazamiento {p['shiftPt'][0]:+.0f}, {p['shiftPt'][1]:+.0f} pt · diff {p['diff']:.3f}"
                    for p in vis.get("pages", [])
                    if "shiftPt" in p
                }
        if panels:
            try:
                result["sheets"] = build_sheets(out_dir, job["orig"], panels, sheets.get("dpi", 70), sheets.get("maxPages", 12), sheets.get("overlay", False), notes)
            except Exception as err:
                result["sheetsError"] = str(err)
    result["seconds"] = round(time.time() - t0, 2)
    with open(os.path.join(out_dir, "metrics.json"), "w", encoding="utf8", newline="\n") as f:
        json.dump(result, f, ensure_ascii=False, indent=1)
    return result


# ── Tools ──────────────────────────────────────────────────────────────────────────────────────────────────────────


def selftest(gt_docx, gt_pdf):
    """The resolver is right when the .docx styles agree with Word's own rendering of the same file."""
    model = DocxModel(gt_docx)
    rendered = pdf_analysis(gt_pdf)
    body = [w for w in docx_style_words(model)]
    header_words = [dict(w, para=-1) for p in model.header_paragraphs for w in p["words"]]
    res = style_agreement(body + header_words, rendered["words"])
    sim = text_similarity(model.text(headers=True), rendered["text"])
    print(json.dumps({"docx": os.path.basename(gt_docx), "styleVsWordRender": res, "textSim": sim}, ensure_ascii=False, indent=1))


def dump(path):
    model = DocxModel(path)
    print(json.dumps(model.structure(), ensure_ascii=False))
    for p in model.paragraphs:
        if not p["text"].strip():
            continue
        styles = Counter(f"{w['style']['family']} {w['style']['size']:g}{' B' if w['style']['bold'] else ''}{' I' if w['style']['italic'] else ''} #{w['style']['color']}" for w in p["words"])
        tags = [p["where"], p["align"]]
        if p["heading"]:
            tags.append(f"H{p['heading']}")
        if p["list"]:
            tags.append(f"{p['list']['kind']}{p['list']['level']}")
        print(f"[{p['index']}] ({', '.join(tags)}) {styles.most_common(2)} | {p['text'][:110]!r}")
    for t in model.tables:
        print(f"table {t['index']}: {dims_label(t)} merged={t['merged']}")
        for row in t["grid"][:6]:
            print("   ", " | ".join(c[:18] for c in row))


def main():
    if hasattr(sys.stdout, "reconfigure"):
        sys.stdout.reconfigure(encoding="utf-8")
    args = sys.argv[1:]
    if not args:
        print(__doc__)
        sys.exit(2)
    if args[0] == "--selftest":
        selftest(args[1], args[2])
    elif args[0] == "--dump":
        dump(args[1])
    else:
        with open(args[0], encoding="utf8") as f:
            job = json.load(f)
        res = run_job(job)
        brief = {k: (v.get("score") if isinstance(v, dict) and "score" in v else None) for k, v in res.items() if k in ("ours", "word")}
        print(json.dumps({"id": res["id"], "scores": brief, "seconds": res["seconds"]}))


if __name__ == "__main__":
    main()
